"""Build outputs/JnJ_Data_Raw_to_Decision.xlsx: the raw data, exactly as downloaded, traced through every step.

Tabs:
  README             how to read the workbook
  Cleaning_Steps     each cleaning step with rows and dollars before/after, plus the column dictionary
  Excluded_Payees    the 100 largest excluded payees and why they were excluded
  Raw_Data_Traced    all 344,504 raw rows (or a slice of them per part when split: build_data_workbook.py <csv> 5), all 32 original columns untouched, plus trace columns:
                     step 1 status + exclusion reason + parsed amount, step 2 supplier segment,
                     step 3 category quadrant, step 4 supplier decision + rule
"""
import sqlite3
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

RAW = Path(sys.argv[1] if len(sys.argv) > 1 else "data/raw/purchase-order-data-2012-2015-.csv")
# Optional 2nd argument: split Raw_Data_Traced into N workbooks (each small enough to copy or email)
PARTS = int(sys.argv[2]) if len(sys.argv) > 2 else 1
DB = Path("data/procurement.db")
OUT = Path("outputs/JnJ_Data_Raw_to_Decision.xlsx")
F = "Arial"
HDR = PatternFill("solid", fgColor="1F3864")
TRACE_HDR = PatternFill("solid", fgColor="C55A11")


def cell(ws, value, bold=False, color="000000", fill=None, fmt=None, size=10, italic=False):
    c = WriteOnlyCell(ws, value=value)
    c.font = Font(name=F, bold=bold, color=color, size=size, italic=italic)
    if fill:
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if fmt:
        c.number_format = fmt
    return c


def table(ws, rows, header, widths=None):
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
    ws.append([cell(ws, h, bold=True, color="FFFFFF", fill=HDR) for h in header])
    for r in rows:
        ws.append([cell(ws, v) if not isinstance(v, tuple) else cell(ws, v[0], fmt=v[1]) for v in r])


raw = pd.read_csv(RAW, dtype=str, keep_default_na=False)
with sqlite3.connect(DB) as con:
    st = pd.read_sql("SELECT * FROM raw_line_status ORDER BY raw_row", con)
    seg = pd.read_sql("SELECT supplier_code, segment FROM supplier_segments", con)
    dec = pd.read_sql("SELECT supplier_code, decision, rule FROM supplier_decisions", con)
quad = pd.read_csv("outputs/step3_kraljic_categories.csv")[["category", "quadrant"]]
assert len(raw) == len(st), "raw file and recorded status must have the same rows"

kept = st.status.values == "KEPT"
trace = pd.DataFrame({"raw_row": st.raw_row, "status": st.status,
                      "exclusion_reason": st.exclusion_reason.replace("", None),
                      "amount": st.amount_usd_parsed})
trace["supplier_code"] = raw["Supplier Code"].values
trace["family"] = raw["Family Title"].values
trace = trace.merge(seg, on="supplier_code", how="left").merge(dec, on="supplier_code", how="left")
trace = trace.merge(quad.rename(columns={"category": "family"}), on="family", how="left")
trace = trace.sort_values("raw_row").reset_index(drop=True)
for c in ["segment", "decision", "rule", "quadrant"]:
    trace.loc[~kept, c] = None
trace.loc[kept & trace.quadrant.isna(), "quadrant"] = "Unclassified"

wb = Workbook(write_only=True)

# ---------------------------------------------------------------- README
ws = wb.create_sheet("README")
ws.column_dimensions["A"].width = 130
for line, bold in [
    ("J&J Supplier Rationalization — raw data traced through every step", True),
    ("Source: State of California DGS, Purchase Order Data FY2012-15, https://catalog.data.gov/dataset/purchase-order-data "
     "(public, no use restrictions). Downloaded Oct 2026.", False),
    ("", False),
    ("Tabs", True),
    ("Cleaning_Steps — each cleaning step in order, with rows and dollars removed, and how every column was parsed.", False),
    ("Excluded_Payees — the 100 largest payees removed in cleaning, and the rule that removed them.", False),
    ("Raw_Data_Traced — every raw row exactly as downloaded (blue headers = original 32 columns, untouched), followed by "
     "orange trace columns:", False),
    ("     Step 1 status (KEPT / EXCLUDED) · Step 1 exclusion reason · Step 1 parsed amount ($) · Step 2 supplier segment · "
     "Step 3 category quadrant · Step 4 supplier decision · Step 4 rule", False),
    ("     Filter 'Step 1 status' = EXCLUDED to see exactly what was cleaned out. Filter 'Step 4 decision' to see the "
     "purchase lines behind each decision.", False),
    ("", False),
    ("Rules R1–R8 for step 4 and the full method are in the outputs workbook (JnJ_Supplier_Rationalization_Outputs.xlsx) "
     "and the GitHub repo.", False),
]:
    ws.append([cell(ws, line, bold=bold, size=14 if line.startswith("J&J") else 10)])

# ---------------------------------------------------------------- Cleaning steps
ws = wb.create_sheet("Cleaning_Steps")
amt = st.amount_usd_parsed
steps = [("0", "Read raw CSV exactly as downloaded", "", len(raw), amt.sum())]
remaining = pd.Series(True, index=st.index)
for code, label, rule in [
    ("exempt_method", "Exclude procurement-exempt purchases",
     "Acquisition Method starts with 'Services are specifically exempt' (by statute or by policy): not negotiable spend"),
    ("public_body_or_health_plan", "Exclude public bodies and health plans",
     "Supplier Name matches: county, city of, state agencies, university, courts, commissions, Medi-Cal health plans "
     "(Health Net, L.A. Care, CalOptima ...), Delta Dental, school/water/transit districts, or 'Unknown'"),
    ("missing_or_zero_amount", "Exclude missing or zero amounts", "Total Price blank, unparseable or $0"),
]:
    m = st.exclusion_reason == code
    remaining &= ~m
    steps.append((str(len(steps)), label, rule, -int(m.sum()), -amt[m].sum()))
steps.append(("=", "Addressable spend (kept)", "Loaded into SQLite table spend_lines", int(remaining.sum()),
              amt[remaining].sum()))
ws.append([cell(ws, "Cleaning steps (step 1)", bold=True, size=12)])
table(ws, [(a, b, c, (d, "#,##0;(#,##0)"), (e / 1e6, '$#,##0.0;($#,##0.0)')) for a, b, c, d, e in steps],
      ["#", "Step", "Rule", "Rows", "Spend ($M)"], [6, 40, 110, 12, 14])
ws.append([])
ws.append([cell(ws, "Parsing applied before the rules", bold=True, size=12)])
table(ws, [
    ("Total Price", "Text like '$6,080.26 ' → number (strip $, commas, spaces)", "amount_usd"),
    ("Unit Price", "Same as Total Price", "unit_price"),
    ("Quantity", "Text → number", "quantity"),
    ("Creation Date", "Text m/d/yyyy → date", "creation_date"),
    ("Supplier Name", "Trim spaces; blank → 'Unknown'", "supplier_name"),
    ("All column names", "Lowercased, spaces/punctuation → underscore", "e.g. 'Supplier Code' → supplier_code"),
], ["Raw column", "Transformation", "Cleaned column"])
ws.append([])
ws.append([cell(ws, "Raw columns not carried into the analysis table", bold=True, size=12)])
table(ws, [(c,) for c in ["Purchase Date", "LPA Number", "Requisition Number", "Sub-Acquisition Type",
                          "Sub-Acquisition Method", "Item Description", "Classification Codes", "Location",
                          "REMOVE AMERISOURCE"]], ["Column (kept in the raw tab, unused in analysis)"])
ws.append([])
ws.append([cell(ws, "What happened to the kept rows in later steps", bold=True, size=12)])
k = trace[kept]
later = []
for col, label in [("segment", "Step 2 supplier segment"), ("quadrant", "Step 3 category quadrant"),
                   ("decision", "Step 4 supplier decision")]:
    g = k.groupby(col).amount.agg(["count", "sum"])
    for name, r in g.iterrows():
        later.append((label, name, (int(r["count"]), "#,##0"), (r["sum"] / 1e6, '$#,##0.0')))
table(ws, later, ["Step", "Value", "Rows", "Spend ($M)"])

# ---------------------------------------------------------------- Excluded payees
ws = wb.create_sheet("Excluded_Payees")
ex = trace[~kept].assign(name=raw.loc[~kept, "Supplier Name"].values)
top = (ex.groupby(["name", "exclusion_reason"]).amount.agg(["count", "sum"]).reset_index()
       .sort_values("sum", ascending=False).head(100))
table(ws, [(r["name"], r["exclusion_reason"], (int(r["count"]), "#,##0"), (r["sum"] / 1e6, '$#,##0.0'))
           for _, r in top.iterrows()], ["Payee (as in raw data)", "Exclusion reason", "Rows", "Spend ($M)"],
      [60, 30, 10, 14])

# ---------------------------------------------------------------- Raw data traced
trace_cols = ["Step 1 status", "Step 1 exclusion reason", "Step 1 parsed amount ($)", "Step 2 supplier segment",
              "Step 3 category quadrant", "Step 4 supplier decision", "Step 4 rule"]
tvals = trace[["status", "exclusion_reason", "amount", "segment", "quadrant", "decision", "rule"]].astype(object)
tvals = tvals.where(tvals.notna(), None).values.tolist()
rows = list(raw.itertuples(index=False, name=None))
n = len(rows)
bounds = [round(n * k / PARTS) for k in range(PARTS + 1)]


def write_raw(book, lo, hi):
    ws = book.create_sheet("Raw_Data_Traced")
    ws.freeze_panes = "A2"
    ws.append([cell(ws, "Raw row #", bold=True, color="FFFFFF", fill=HDR)]
              + [cell(ws, c, bold=True, color="FFFFFF", fill=HDR) for c in raw.columns]
              + [cell(ws, c, bold=True, color="FFFFFF", fill=TRACE_HDR) for c in trace_cols])
    for i in range(lo, hi):
        tv = list(tvals[i])
        if tv[2] is not None and tv[2] != tv[2]:  # NaN amount
            tv[2] = None
        ws.append([i + 1, *rows[i], *tv])


OUT.parent.mkdir(exist_ok=True)
if PARTS == 1:
    write_raw(wb, 0, n)
    wb.save(OUT)
    print("saved", OUT, f"{OUT.stat().st_size/1e6:.1f} MB")
else:
    for k in range(PARTS):
        lo, hi = bounds[k], bounds[k + 1]
        if k == 0:
            book = wb  # part 1 also carries README, Cleaning_Steps and Excluded_Payees
        else:
            book = Workbook(write_only=True)
            w = book.create_sheet("README")
            w.column_dimensions["A"].width = 120
            w.append([cell(w, f"J&J Supplier Rationalization — raw data traced, part {k+1} of {PARTS}", bold=True, size=14)])
            w.append([cell(w, f"Raw rows {lo+1:,} to {hi:,} of {n:,}. Tabs README, Cleaning_Steps and Excluded_Payees are in part 1.")])
        write_raw(book, lo, hi)
        out = OUT.with_name(f"{OUT.stem}_part{k+1}of{PARTS}.xlsx")
        book.save(out)
        print("saved", out, f"rows {lo+1:,}-{hi:,}", f"{out.stat().st_size/1e6:.1f} MB", flush=True)
