"""Build outputs/JnJ_Supplier_Rationalization_Outputs.xlsx — one tab per completed step."""
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path("outputs/JnJ_Supplier_Rationalization_Outputs.xlsx")
F = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
USD_M = '$#,##0.0;($#,##0.0);"-"'
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'

wb = Workbook()


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(italic=True, color="595959")


def header(ws, row, cols, start_col=1):
    for i, c in enumerate(cols):
        cell = ws.cell(row=row, column=start_col + i, value=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_df(ws, df, row, fmts=None, start_col=1):
    """Write a DataFrame with a header row; fmts maps column name -> number format."""
    header(ws, row, list(df.columns), start_col)
    for r, rec in enumerate(df.itertuples(index=False), start=row + 1):
        for c, v in enumerate(rec, start=start_col):
            cell = ws.cell(row=r, column=c, value=None if pd.isna(v) else v)
            cell.font = font()
            col = df.columns[c - start_col]
            if fmts and col in fmts:
                cell.number_format = fmts[col]
    return row + len(df)


def widths(ws, w):
    for i, x in enumerate(w, start=1):
        ws.column_dimensions[get_column_letter(i)].width = x


# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
title(ws, "J&J Supplier Rationalization — outputs by step",
      "Real data: State of California purchase orders FY2012-15 (data.gov); J&J figures from J&J's public reports")
rows = [
    ("Tab", "Step", "Question it answers", "Method"),
    ("Files_and_Sources", "All", "Where every number comes from", "Data sources, code, SQL, outputs with links"),
    ("0_Business_Case", "0 Business case", "Is this worth doing for J&J?", "Top-down savings model with live formulas"),
    ("1_Data_Cleaning", "1 Load real data", "What spend is addressable?", "Exclusion rules with reason codes"),
    ("2_Segments", "2 Spend cube + Pareto/ABC", "Where does the spend sit?", "SQL window functions"),
    ("2_Fragmented_Categories", "2 Spend cube", "Which categories have too many suppliers?", "SQL category cube"),
    ("2_Supplier_List", "2 Spend cube", "Segment of every supplier", "SQL (16,933 rows)"),
    ("3_Kraljic", "3 Kraljic matrix", "How should each category be handled?", "Value vs supply-risk scoring"),
    ("4_Decision_Summary", "4 Keep / consolidate / exit", "Which suppliers stay and which go?", "Duplicate detection (ML) + scorecard + rules"),
    ("4_Supplier_Decisions", "4 Keep / consolidate / exit", "Decision and reason for every supplier", "16,933 rows, filterable"),
    ("4_Concentration", "4c Supply-risk check", "Does consolidation concentrate spend?", "Top-supplier share per category, before vs after"),
    ("4_Duplicate_Check", "4d Duplicate precision", "How often is a matched pair really one company?", "70 accepted pairs judged by hand"),
    ("5_Savings_Real_List", "5 Savings on real list", "What is it worth?", "Measured price gap + scenarios, live formulas"),
    ("5b_Retender_Sensitivity", "5b Sensitivity", "How much rests on the re-tender assumption?", "Savings vs re-tender rate, break-even, live formulas"),
    ("5c_Value_vs_Cuts", "5c Small cut, big impact", "Does cutting more suppliers pay?", "Savings and one-time cost per tier at the same rates"),
    ("8_Phase2", "8 Phase 2", "How could J&J go further from ~39,000?", "Category target count with supply-risk floors"),
    ("9_JnJ_Timeline", "9 Back-test", "Does the model explain J&J's 80,000 to ~39,000?", "Reported counts, Kenvue spin-off estimate, tiers applied"),
    ("SQL_Code", "2", "The SQL behind the segmentation", "Full query text"),
]
for r, rec in enumerate(rows, start=4):
    for c, v in enumerate(rec, start=1):
        ws.cell(row=r, column=c, value=v).font = font(bold=(r == 4))
header(ws, 4, list(rows[0]))
ws["A23"] = "Colour code: blue = input you can change; black = formula; yellow fill = key assumption."
ws["A23"].font = font(italic=True)
widths(ws, [26, 26, 40, 44])

# ---------------------------------------------------------------- Step 0
ws = wb.create_sheet("0_Business_Case")
title(ws, "Step 0 — Business case (J&J, base year)",
      "Inputs in blue drive every number below. $ in millions unless stated.")
inputs = [
    ("Addressable spend ($M)", 35000, USD, "J&J careers posting: ~$35B across ~35,000 suppliers"),
    ("Active suppliers", 29000, NUM, "Step-0 estimate only (2024 Health for Humanity Report: 29,000+). Superseded: step 5 scales to 80,000"),
    ("Annual cost to maintain one supplier ($M)", 0.002, '$#,##0.000', "Assumption: $2,000 per supplier per year"),
    ("Program team + analytics ($M)", 15, USD_M, "Assumption"),
    ("Contracting and legal ($M)", 5, USD_M, "Assumption"),
    ("Contingency", 0.15, PCT, "Assumption"),
    ("Discount rate", 0.08, PCT, "Assumption"),
    ("Savings ramp year 1", 0.30, PCT, "Assumption"),
    ("Savings ramp year 2", 0.70, PCT, "Assumption"),
    ("Savings ramp year 3", 1.00, PCT, "Assumption"),
]
header(ws, 4, ["Input", "Value", "Source / note"])
for i, (lab, v, fmt, src) in enumerate(inputs, start=5):
    ws.cell(row=i, column=1, value=lab).font = font()
    c = ws.cell(row=i, column=2, value=v)
    c.font, c.number_format, c.fill = font(BLUE), fmt, KEY_FILL
    ws.cell(row=i, column=3, value=src).font = font(italic=True, color="595959")
SPEND, SUP, COST, PROG, LEGAL, CONT, DISC, R1, R2, R3 = [f"$B${i}" for i in range(5, 15)]

r0 = 17
ws.cell(row=r0 - 1, column=1, value="Segments (shares validated on real data in step 2)").font = font(bold=True)
cols = ["Segment", "Share of suppliers", "Share of spend", "Target cut", "Price savings rate",
        "Transition cost per exit ($M)", "Suppliers", "Spend ($M)", "Exits", "Suppliers after",
        "Spend addressed ($M)", "Price savings ($M)", "Transition cost ($M)"]
header(ws, r0, cols)
segs = [("A Strategic", 0.02, 0.70, 0.00, 0.005, 0.0),
        ("B Core", 0.13, 0.22, 0.25, 0.05, 0.025),
        ("C Tail", 0.85, 0.08, 0.60, 0.08, 0.0005)]
for i, (name, ss, sp, cut, rate, tr) in enumerate(segs, start=r0 + 1):
    ws.cell(row=i, column=1, value=name).font = font()
    for col, v, fmt in [(2, ss, PCT), (3, sp, PCT), (4, cut, PCT), (5, rate, PCT), (6, tr, '$#,##0.0000')]:
        c = ws.cell(row=i, column=col, value=v)
        c.font, c.number_format = font(BLUE), fmt
    f = {
        7: (f"={SUP}*B{i}", NUM), 8: (f"={SPEND}*C{i}", USD), 9: (f"=G{i}*D{i}", NUM), 10: (f"=G{i}-I{i}", NUM),
        11: (f"=IF(D{i}=0,H{i},H{i}*D{i})", USD), 12: (f"=K{i}*E{i}", USD_M), 13: (f"=I{i}*F{i}", USD_M),
    }
    for col, (formula, fmt) in f.items():
        c = ws.cell(row=i, column=col, value=formula)
        c.font, c.number_format = font(), fmt
tr_ = r0 + 4
ws.cell(row=tr_, column=1, value="Total").font = font(bold=True)
for col, fmt in [(7, NUM), (8, USD), (9, NUM), (10, NUM), (11, USD), (12, USD_M), (13, USD_M)]:
    L = get_column_letter(col)
    c = ws.cell(row=tr_, column=col, value=f"=SUM({L}{r0+1}:{L}{r0+3})")
    c.font, c.number_format = font(bold=True), fmt
ws.cell(row=tr_ + 1, column=1, value="Supplier reduction").font = font()
c = ws.cell(row=tr_ + 1, column=9, value=f"=I{tr_}/G{tr_}")
c.font, c.number_format = font(), PCT

r = tr_ + 3
summary = [
    ("Price savings ($M/yr)", f"=L{tr_}"),
    ("Process savings: fewer suppliers to maintain ($M/yr)", f"=I{tr_}*{COST}"),
    ("Run-rate savings ($M/yr)", f"=B{r}+B{r+1}"),
    ("Run-rate as % of addressable spend", f"=B{r+2}/{SPEND}"),
    ("One-time cost ($M)", f"=({PROG}+{LEGAL}+M{tr_})*(1+{CONT})"),
]
for i, (lab, formula) in enumerate(summary):
    ws.cell(row=r + i, column=1, value=lab).font = font(bold=(i in (2, 4)))
    c = ws.cell(row=r + i, column=2, value=formula)
    c.font, c.number_format = font(bold=(i in (2, 4))), (PCT if i == 3 else USD_M)
RUN, OT = f"$B${r+2}", f"$B${r+4}"

rs = r + 7
ws.cell(row=rs - 1, column=1, value="Scenarios").font = font(bold=True)
header(ws, rs, ["Scenario", "Savings vs base", "Cost vs base", "Run-rate ($M/yr)", "Year 1 ($M)", "Year 2 ($M)",
                "Year 3 ($M)", "3-yr savings ($M)", "One-time cost ($M)", "3-yr NPV ($M)", "ROI"])
for i, (name, a, b) in enumerate([("Low", 0.5, 1.25), ("Base", 1.0, 1.0), ("High", 1.3, 0.9)], start=rs + 1):
    ws.cell(row=i, column=1, value=name).font = font()
    for col, v in [(2, a), (3, b)]:
        c = ws.cell(row=i, column=col, value=v)
        c.font, c.number_format = font(BLUE), PCT
    f = {4: f"={RUN}*B{i}", 5: f"=D{i}*{R1}", 6: f"=D{i}*{R2}", 7: f"=D{i}*{R3}", 8: f"=SUM(E{i}:G{i})",
         9: f"={OT}*C{i}", 10: f"=-I{i}+E{i}/(1+{DISC})+F{i}/(1+{DISC})^2+G{i}/(1+{DISC})^3",
         11: f"=(H{i}-I{i})/I{i}"}
    for col, formula in f.items():
        c = ws.cell(row=i, column=col, value=formula)
        c.font, c.number_format = font(), ('0.0"x"' if col == 11 else USD_M)
widths(ws, [44, 14, 14, 12, 12, 14, 12, 12, 12, 14, 14, 14, 14])

# ---------------------------------------------------------------- data
con = sqlite3.connect("data/procurement.db")

# ---------------------------------------------------------------- Step 1
ws = wb.create_sheet("1_Data_Cleaning")
title(ws, "Step 1 — Load and clean real data",
      "Source: State of California DGS, Purchase Order Data FY2012-15, https://catalog.data.gov/dataset/purchase-order-data")
ex = pd.read_sql("SELECT exclusion_reason AS reason, COUNT(*) AS lines, SUM(amount_usd)/1e6 AS spend_m "
                 "FROM excluded_lines GROUP BY 1", con)
kept = pd.read_sql("SELECT COUNT(*) AS lines, SUM(amount_usd)/1e6 AS spend_m, "
                   "COUNT(DISTINCT supplier_code) AS suppliers FROM spend_lines", con).iloc[0]
label = {"exempt_method": "Excluded: procurement-exempt methods",
         "public_body_or_health_plan": "Excluded: public bodies and health plans (pass-through)",
         "missing_or_zero_amount": "Excluded: missing or zero amount"}
header(ws, 4, ["Stage", "Lines", "Spend ($M)"])
rows = [(label[x.reason], x.lines, x.spend_m) for x in ex.itertuples()]
for i, (lab, n, s) in enumerate(rows, start=6):
    ws.cell(row=i, column=1, value=lab).font = font()
    ws.cell(row=i, column=2, value=int(n)).number_format = NUM
    ws.cell(row=i, column=3, value=float(s)).number_format = USD_M
k = 6 + len(rows)
ws.cell(row=k, column=1, value="Addressable (kept)").font = font(bold=True)
ws.cell(row=k, column=2, value=int(kept.lines)).number_format = NUM
ws.cell(row=k, column=3, value=float(kept.spend_m)).number_format = USD_M
ws.cell(row=5, column=1, value="Raw file").font = font(bold=True)
for col in (2, 3):
    L = get_column_letter(col)
    c = ws.cell(row=5, column=col, value=f"=SUM({L}6:{L}{k})")
    c.font, c.number_format = font(bold=True), (NUM if col == 2 else USD_M)
ws.cell(row=k + 2, column=1, value="Addressable suppliers").font = font()
ws.cell(row=k + 2, column=2, value=int(kept.suppliers)).number_format = NUM
ws.cell(row=k + 3, column=1, value="Addressable share of raw spend").font = font()
c = ws.cell(row=k + 3, column=2, value=f"=C{k}/C5")
c.number_format = PCT
widths(ws, [52, 14, 16])

# ---------------------------------------------------------------- Step 2
ws = wb.create_sheet("2_Segments")
title(ws, "Step 2 — Where does the spend sit? (Pareto / ABC)",
      "A = suppliers making up the first 70% of spend; B = next 22%; C = long tail. 3-year spend.")
seg = pd.read_sql("""SELECT segment AS Segment, COUNT(*) AS Suppliers, SUM(spend_usd)/1e6 AS "Spend ($M)",
    SUM(inactive_18m) AS "Inactive 18+ months", SUM(one_time) AS "One-time (1 PO)",
    SUM(CASE WHEN inactive_18m=1 OR one_time=1 THEN 1 ELSE 0 END) AS "Inactive or one-time",
    SUM(small_or_diverse) AS "Small or diverse"
    FROM supplier_segments GROUP BY 1 ORDER BY 1""", con)
end = write_df(ws, seg, 4, {"Suppliers": NUM, "Spend ($M)": USD_M, "Inactive 18+ months": NUM,
                            "One-time (1 PO)": NUM, "Inactive or one-time": NUM, "Small or diverse": NUM})
t = end + 1
ws.cell(row=t, column=1, value="Total").font = font(bold=True)
for col in range(2, 8):
    L = get_column_letter(col)
    c = ws.cell(row=t, column=col, value=f"=SUM({L}5:{L}{end})")
    c.font, c.number_format = font(bold=True), (USD_M if col == 3 else NUM)
header(ws, 4, ["Segment", "Suppliers", "Spend ($M)", "Inactive 18+ months", "One-time (1 PO)",
               "Inactive or one-time", "Small or diverse", "% of suppliers", "% of spend", "% of segment inactive or one-time"])
for i in range(5, end + 1):
    for col, formula in [(8, f"=B{i}/B${t}"), (9, f"=C{i}/C${t}"), (10, f"=F{i}/B{i}")]:
        c = ws.cell(row=i, column=col, value=formula)
        c.font, c.number_format = font(), PCT

p = t + 3
ws.cell(row=p - 1, column=1, value="Pareto curve (real data) vs business case assumption").font = font(bold=True)
curve = pd.read_csv("outputs/step2_pareto_curve.csv")
curve.columns = ["Top % of suppliers", "% of spend (real)"]
curve["Top % of suppliers"] /= 100
curve["% of spend (real)"] /= 100
write_df(ws, curve, p, {"Top % of suppliers": PCT, "% of spend (real)": PCT})
ws.cell(row=p, column=4, value="Business case assumption").font = Font(name=F, bold=True, color="FFFFFF")
ws.cell(row=p, column=4).fill = HDR_FILL
ws.cell(row=p + 3, column=4, value="2% of suppliers = 70% of spend").font = font(italic=True)
widths(ws, [16, 12, 14, 16, 14, 16, 14, 14, 12, 18])

ws = wb.create_sheet("2_Fragmented_Categories")
title(ws, "Step 2 — Most fragmented categories (most suppliers)",
      "Many suppliers with little spend each = first consolidation targets")
frag = pd.read_csv("outputs/step2_fragmented_categories.csv")
frag.columns = ["Category", "Suppliers", "Tail suppliers", "Spend ($M)", "Spend per supplier ($K)",
                "% of spend with tail suppliers"]
frag["% of spend with tail suppliers"] /= 100
write_df(ws, frag, 4, {"Suppliers": NUM, "Tail suppliers": NUM, "Spend ($M)": USD_M,
                       "Spend per supplier ($K)": USD, "% of spend with tail suppliers": PCT})
widths(ws, [60, 12, 14, 12, 18, 20])

ws = wb.create_sheet("2_Supplier_List")
title(ws, "Step 2 — Every supplier with its segment (16,933 real suppliers)",
      "Spend is 3-year addressable spend. Inactive = no PO after 2013-12-31. One-time = one PO in 3 years.")
sl = pd.read_sql("""SELECT spend_rank AS Rank, supplier_code AS "Supplier code", supplier_name AS Supplier,
    segment AS Segment, spend_usd AS "Spend ($)", po_count AS POs, departments AS Departments,
    categories AS Categories, substr(last_po_date,1,10) AS "Last PO", inactive_18m AS "Inactive 18m",
    one_time AS "One-time", small_or_diverse AS "Small or diverse", cum_share AS "Cumulative % of spend"
    FROM supplier_segments ORDER BY spend_rank""", con)
write_df(ws, sl, 4, {"Spend ($)": USD, "POs": NUM, "Cumulative % of spend": PCT})
ws.freeze_panes = "D5"
ws.auto_filter.ref = f"A4:{get_column_letter(len(sl.columns))}{4 + len(sl)}"
widths(ws, [8, 13, 46, 13, 15, 8, 12, 11, 12, 11, 10, 14, 14])

# ---------------------------------------------------------------- Step 3
ws = wb.create_sheet("3_Kraljic")
title(ws, "Step 3 — How should each category be handled? (Kraljic matrix)",
      "Value = category spend (top categories = 80% of spend). Risk = non-competitive share + top-3 concentration "
      "+ thin supply base, ranked within size bands.")
ks = pd.read_csv("outputs/step3_kraljic_summary.csv")
ks = ks[["quadrant", "categories", "spend_b", "suppliers"]]
ks.columns = ["Quadrant", "Categories", "Spend ($B)", "Supplier relationships"]
end = write_df(ws, ks, 4, {"Categories": NUM, "Spend ($B)": '$#,##0.00', "Supplier relationships": NUM})
ws.cell(row=4, column=5, value="% of spend")
header(ws, 4, ["Quadrant", "Categories", "Spend ($B)", "Supplier relationships", "% of spend"])
for i in range(5, end + 1):
    c = ws.cell(row=i, column=5, value=f"=C{i}/SUM(C$5:C${end})")
    c.font, c.number_format = font(), PCT
kc = pd.read_csv("outputs/step3_kraljic_categories.csv")
kc = kc[["category", "segment_title", "quadrant", "agreement", "spend_usd", "suppliers", "active_suppliers",
         "pct_noncompetitive", "pct_top3", "risk_score"]]
kc.columns = ["Category", "UNSPSC segment", "Quadrant", "Agreement type", "Spend ($)", "Suppliers",
              "Active suppliers", "% non-competitive", "% with top 3 suppliers", "Risk score (0-1)"]
kc["% non-competitive"] /= 100
kc["% with top 3 suppliers"] /= 100
s = end + 3
ws.cell(row=s - 1, column=1, value="All 409 categories").font = font(bold=True)
write_df(ws, kc, s, {"Spend ($)": USD, "Suppliers": NUM, "Active suppliers": NUM,
                     "% non-competitive": PCT, "% with top 3 suppliers": PCT, "Risk score (0-1)": "0.00"})
ws.auto_filter.ref = f"A{s}:J{s + len(kc)}"
widths(ws, [52, 40, 13, 46, 15, 11, 11, 14, 14, 12])

# ---------------------------------------------------------------- Step 4
if Path("outputs/step4_supplier_decisions.csv").exists():
    ws = wb.create_sheet("4_Decision_Summary")
    title(ws, "Step 4 — Which suppliers stay and which go?",
          "Rules applied in order; the first rule that matches sets the decision. Duplicates found by TF-IDF "
          "name matching (cosine >= 0.88).")
    br = pd.read_csv("outputs/step4_decisions_by_rule.csv")
    br.columns = ["Rule", "Decision", "Reason", "Suppliers", "Spend 3 yrs ($M)"]
    end = write_df(ws, br, 4, {"Suppliers": NUM, "Spend 3 yrs ($M)": USD_M})
    t = end + 1
    ws.cell(row=t, column=1, value="Total").font = font(bold=True)
    for col, fmt in [(4, NUM), (5, USD_M)]:
        L = get_column_letter(col)
        c = ws.cell(row=t, column=col, value=f"=SUM({L}5:{L}{end})")
        c.font, c.number_format = font(bold=True), fmt
    h = pd.read_csv("outputs/step4_headline.csv")
    r = t + 3
    ws.cell(row=r - 1, column=1, value="Headline").font = font(bold=True)
    header(ws, r, ["Metric", "Value"])
    for i, rec in enumerate(h.itertuples(index=False), start=r + 1):
        ws.cell(row=i, column=1, value=rec.metric).font = font()
        c = ws.cell(row=i, column=2, value=float(rec.value))
        c.number_format = PCT if rec.value < 1 else (USD_M if "$M" in rec.metric else NUM)
    seg = pd.read_csv("outputs/step4_decisions_by_segment.csv")
    r = r + len(h) + 3
    ws.cell(row=r - 1, column=1, value="Decisions by segment").font = font(bold=True)
    write_df(ws, seg, r, {c: NUM for c in seg.columns[1:]})
    widths(ws, [66, 14, 70, 12, 16])

    ws = wb.create_sheet("4_Supplier_Decisions")
    title(ws, "Step 4 — Decision and reason code for every supplier (16,933 real suppliers)",
          "Filter by Decision or Rule. Score = 0-100 scorecard within the supplier's primary category.")
    sd = pd.read_csv("outputs/step4_supplier_decisions.csv")
    sd.columns = ["Supplier code", "Supplier", "Segment", "Primary category", "Kraljic quadrant", "Spend 3 yrs ($)",
                  "POs", "Departments", "Last PO", "Small or diverse", "Score", "Rule", "Decision", "Reason",
                  "Merge into (master)", "Agreement if kept"]
    sd["Last PO"] = sd["Last PO"].astype(str).str[:10]
    write_df(ws, sd, 4, {"Spend 3 yrs ($)": USD, "POs": NUM, "Score": "0.0"})
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(sd.columns))}{4 + len(sd)}"
    widths(ws, [12, 42, 12, 36, 13, 15, 7, 11, 11, 10, 8, 6, 13, 60, 36, 34])

# ---------------------------------------------------------------- Step 5
if Path("outputs/step5_model_inputs.csv").exists():
    ws = wb.create_sheet("5_Savings_Real_List")
    title(ws, "Step 5 — What is it worth? Savings on the real keep / exit list",
          "Blue = input. Data inputs come from steps 4-5; rates marked 'measured' come from the price-gap test. "
          "All $ in millions per year.")
    inp = dict(pd.read_csv("outputs/step5_model_inputs.csv").values)
    tiers = pd.read_csv("outputs/step5_jnj_supplier_tiers.csv")
    gap = pd.read_csv("outputs/step5_measured_price_gap.csv").set_index("quadrant")

    header(ws, 4, ["Data input (from the real list)", "Value", "Source"])
    data_rows = [
        ("CA addressable spend per year ($M)", inp["CA addressable spend per year ($M)"], USD_M, "Step 1, 3-yr / 3"),
        ("Re-routed spend, Leverage ($M/yr)", inp["Re-routed spend per year, base plan, Leverage ($M)"], USD_M, "Step 4: R5 + R8 in Leverage"),
        ("Re-routed spend, Non-critical ($M/yr)", inp["Re-routed spend per year, base plan, Non-critical ($M)"], USD_M, "Step 4: R5 + R8 in Non-critical"),
        ("Re-routed spend, Bottleneck ($M/yr)", inp["Re-routed spend per year, base plan, Bottleneck ($M)"], USD_M, "Step 4: R5 only (no consolidation in base plan)"),
        ("Re-routed spend, Strategic ($M/yr)", inp["Re-routed spend per year, base plan, Strategic ($M)"], USD_M, "Step 4: R5 only (no consolidation in base plan)"),
        ("Kept spend, Leverage categories ($M/yr)", inp["Kept spend per year, Leverage categories ($M)"], USD_M, "Step 4 KEEP"),
        ("CA suppliers active in FY2014-15", inp["CA suppliers active in FY2014-2015"], NUM, "One year of active suppliers, like J&J's annual count"),
        ("Measured price gap, Leverage", gap.loc["Leverage", "rate_used_high"], PCT,
         f"Price-gap test: median of {int(gap.loc['Leverage', 'items'])} goods items"),
        ("Measured price gap, Non-critical", gap.loc["Non-critical", "rate_used_high"], PCT,
         f"Price-gap test: median of {int(gap.loc['Non-critical', 'items'])} goods items"),
        ("Measured price gap, Bottleneck (exiting suppliers cheaper → 0)", 0, PCT,
         f"Measured {gap.loc['Bottleneck', 'median_gap']:.1%}, floored at 0"),
        ("Measured price gap, Strategic (exiting suppliers cheaper → 0)", 0, PCT,
         f"Measured {gap.loc['Strategic', 'median_gap']:.1%}, floored at 0"),
        ("J&J addressable spend ($M/yr)", 35000, USD, "J&J careers posting"),
        ("J&J suppliers", 80000, NUM, "J&J supplier count as reported (State of Flux case study); latest J&J statement FY2025: ~39,000"),
        ("Program cost as % of spend", 20 / 35000, '0.000%', "Business case: $20M on $35B"),
        ("Transition cost per R8 supplier ($)", 2500, USD, "Assumption"),
        ("Transition cost per R5 supplier ($)", 500, USD, "Assumption"),
        ("Transition cost per R3 supplier ($)", 200, USD, "Assumption"),
        ("Contingency", 0.15, PCT, "Assumption"),
        ("Discount rate", 0.08, PCT, "Assumption"),
        ("Savings goal, J&J ($M/yr)", 300, USD, "Project goal"),
    ]
    ref = {}
    for i, (lab, v, fmt, src) in enumerate(data_rows, start=5):
        ws.cell(row=i, column=1, value=lab).font = font()
        c = ws.cell(row=i, column=2, value=float(v))
        c.font, c.number_format = font(BLUE), fmt
        ws.cell(row=i, column=3, value=src).font = font(italic=True, color="595959")
        ref[lab] = f"$B${i}"
    R = lambda k: ref[k]
    spend_ratio = f"({R('J&J addressable spend ($M/yr)')}/{R('CA addressable spend per year ($M)')})"
    sup_ratio = f"({R('J&J suppliers')}/{R('CA suppliers active in FY2014-15')})"

    s0 = 5 + len(data_rows) + 2
    ws.cell(row=s0 - 1, column=1, value="Scenario levers (change these)").font = font(bold=True)
    header(ws, s0, ["Lever", "Low", "Base", "High", "Basis"])
    levers = [("Share of measured price gap captured", (0, 0.5, 1), PCT, "Data shows the gap; low assumes none is captured"),
              ("Re-tender saving on kept Leverage spend", (0.005, 0.015, 0.03), PCT, "ASSUMPTION — not measurable in this data"),
              ("Cost to maintain one active supplier ($/yr)", (1000, 2000, 3000), USD, "Assumption"),
              ("Duplicates merged (R3), active this year", tuple(int(t) for t in tiers.ca_removed_R3), NUM, "Tier: all scenarios"),
              ("One-time vendors to catalog (R5), active this year", tuple(int(t) for t in tiers.ca_removed_R5), NUM, "Tier: all scenarios"),
              ("Suppliers consolidated (R8), active this year", tuple(int(t) for t in tiers.ca_removed_R8), NUM,
               "Conservative: none · Base: Leverage + Non-critical · Aggressive: all categories")]
    for i, (lab, vals, fmt, basis) in enumerate(levers, start=s0 + 1):
        ws.cell(row=i, column=1, value=lab).font = font()
        for j, v in enumerate(vals, start=2):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.number_format, c.fill = font(BLUE), fmt, KEY_FILL
        ws.cell(row=i, column=5, value=basis).font = font(italic=True, color="595959")
    L_CAP, L_RET, L_COST, L_R3, L_R5, L_R8 = s0 + 1, s0 + 2, s0 + 3, s0 + 4, s0 + 5, s0 + 6

    r0 = L_R8 + 3
    ws.cell(row=r0 - 1, column=1, value="Results").font = font(bold=True)
    header(ws, r0, ["Result ($M per year unless stated)", "Low", "Base", "High", "How it's calculated"])
    price_ca = (f"({R('Re-routed spend, Leverage ($M/yr)')}*{R('Measured price gap, Leverage')}"
                f"+{R('Re-routed spend, Non-critical ($M/yr)')}*{R('Measured price gap, Non-critical')}"
                f"+{R('Re-routed spend, Bottleneck ($M/yr)')}*{R('Measured price gap, Bottleneck (exiting suppliers cheaper → 0)')}"
                f"+{R('Re-routed spend, Strategic ($M/yr)')}*{R('Measured price gap, Strategic (exiting suppliers cheaper → 0)')})")
    trans_ca = lambda c: (f"({c}{L_R8}*{R('Transition cost per R8 supplier ($)')}"
                          f"+{c}{L_R5}*{R('Transition cost per R5 supplier ($)')}"
                          f"+{c}{L_R3}*{R('Transition cost per R3 supplier ($)')})/1000000")
    res_rows = [
        ("CA: price gap on re-routed spend", lambda c: f"={price_ca}*{c}{L_CAP}", USD_M, "Re-routed spend × measured gap × share captured"),
        ("CA: re-tender of kept Leverage suppliers", lambda c: f"={R('Kept spend, Leverage categories ($M/yr)')}*{c}{L_RET}", USD_M, "Kept leverage spend × re-tender %"),
        ("CA: process savings", lambda c: f"=({c}{L_R3}+{c}{L_R5}+{c}{L_R8})*{c}{L_COST}/1000000", USD_M, "Suppliers leaving the managed base × cost per supplier"),
        ("CA: run-rate savings", lambda c: f"=SUM({c}{r0+1}:{c}{r0+3})", USD_M, "Sum"),
        ("CA: savings as % of spend", lambda c: f"={c}{r0+4}/{R('CA addressable spend per year ($M)')}", PCT, ""),
        ("J&J: price gap", lambda c: f"={c}{r0+1}*{spend_ratio}", USD_M, "Scaled by spend"),
        ("J&J: re-tender", lambda c: f"={c}{r0+2}*{spend_ratio}", USD_M, "Scaled by spend"),
        ("J&J: process", lambda c: f"={c}{r0+3}*{sup_ratio}", USD_M, "Scaled by supplier count"),
        ("J&J: run-rate savings", lambda c: f"=SUM({c}{r0+6}:{c}{r0+8})", USD_M, "Sum"),
        ("J&J: measured in data (price gap + process)", lambda c: f"={c}{r0+6}+{c}{r0+8}", USD_M, "Excludes the re-tender assumption"),
        ("J&J: share of savings measured in data", lambda c: f"=IF({c}{r0+9}=0,0,{c}{r0+10}/{c}{r0+9})", PCT, ""),
        ("J&J: one-time cost ($M)", lambda c: f"=({R('Program cost as % of spend')}*{R('J&J addressable spend ($M/yr)')}+{trans_ca(c)}*{sup_ratio})*(1+{R('Contingency')})", USD_M,
         "Program (scaled by spend) + transition (scaled by suppliers) + contingency"),
        ("J&J: 3-yr NPV ($M)", lambda c: f"=-{c}{r0+12}+0.3*{c}{r0+9}/(1+{R('Discount rate')})+0.7*{c}{r0+9}/(1+{R('Discount rate')})^2+{c}{r0+9}/(1+{R('Discount rate')})^3",
         USD_M, "Ramp 30% / 70% / 100%"),
        ("Meets $300M goal?", lambda c: f'=IF({c}{r0+9}>={R("Savings goal, J&J ($M/yr)")},"Yes","No")', None, ""),
        ("Re-tender % needed to reach the goal", lambda c: f"=MAX(0,({R('Savings goal, J&J ($M/yr)')}-{c}{r0+10})/({R('Kept spend, Leverage categories ($M/yr)')}*{spend_ratio}))",
         '0.00%', "What has to be true"),
    ]
    for i, (lab, fn, fmt, how) in enumerate(res_rows, start=r0 + 1):
        bold = lab in ("CA: run-rate savings", "J&J: run-rate savings", "Meets $300M goal?")
        ws.cell(row=i, column=1, value=lab).font = font(bold=bold)
        for col in "BCD":
            c = ws[f"{col}{i}"]
            c.value = fn(col)
            c.font = font(bold=bold)
            if fmt:
                c.number_format = fmt
        ws.cell(row=i, column=5, value=how).font = font(italic=True, color="595959")

    g0 = r0 + len(res_rows) + 3
    ws.cell(row=g0 - 1, column=1, value="Price-gap evidence: goods items bought from both kept and exiting suppliers").font = font(bold=True)
    pg = pd.read_csv("outputs/step5_price_gap_goods.csv").sort_values("exiting_spend", ascending=False)
    pg = pg[["unspsc", "item", "quadrant", "kept_median_price", "exiting_median_price", "gap", "kept_lines",
             "exiting_lines", "exiting_spend"]]
    pg.columns = ["UNSPSC", "Item", "Quadrant", "Kept suppliers: median unit price ($)",
                  "Exiting suppliers: median unit price ($)", "Gap (exiting dearer = +)", "Kept lines",
                  "Exiting lines", "Exiting spend, 3 yrs ($)"]
    write_df(ws, pg, g0, {"Kept suppliers: median unit price ($)": '$#,##0.00',
                          "Exiting suppliers: median unit price ($)": '$#,##0.00',
                          "Gap (exiting dearer = +)": PCT, "Exiting spend, 3 yrs ($)": USD})
    widths(ws, [58, 30, 14, 26, 60, 14, 12, 12, 18])


# ---------------------------------------------------------------- Step 4c / 4d / 5b
if Path("outputs/step4c_concentration_summary.csv").exists():
    ws = wb.create_sheet("4_Concentration")
    title(ws, "Step 4c — Does consolidation add supply risk? Top-supplier share per category, before vs after",
          "Active spend only (inactive R4 suppliers left out of both views); duplicate records folded into their "
          "master; R5/R8 spend moved to kept suppliers in the same category in proportion to their spend.")
    cs = pd.read_csv("outputs/step4c_concentration_summary.csv")
    cs.columns = ["Group", "Categories", "Top supplier >50%: before", "Top supplier >50%: after",
                  "Single supplier: before", "Single supplier: after", "Median top share: before",
                  "Median top share: after", "Median HHI: before", "Median HHI: after",
                  "Spend in >50% categories, 3 yrs ($M): before", "Spend in >50% categories, 3 yrs ($M): after"]
    end = write_df(ws, cs, 4, {**{c: NUM for c in cs.columns[1:6]}, **{c: PCT for c in cs.columns[6:8]},
                                **{c: "0.000" for c in cs.columns[8:10]}, **{c: USD_M for c in cs.columns[10:]}})
    wl = pd.read_csv("outputs/step4c_concentration_watchlist.csv")
    wl.columns = ["Category", "Quadrant", "Spend after, 3 yrs ($)", "Suppliers before", "Suppliers after",
                  "Top share before", "Top share after"]
    r = end + 3
    ws.cell(row=r - 1, column=1, value="Watch-list: categories where the top supplier crosses 50% after "
            "consolidation (add a second-source clause)").font = font(bold=True)
    end = write_df(ws, wl, r, {"Spend after, 3 yrs ($)": USD, "Suppliers before": NUM, "Suppliers after": NUM,
                               "Top share before": PCT, "Top share after": PCT})
    ss = pd.read_csv("outputs/step4c_sole_source_suppliers.csv")
    ss.columns = ["Category", "Supplier code (master)", "Spend, 3 yrs ($)", "Decision", "Quadrant"]
    r = end + 3
    ws.cell(row=r - 1, column=1, value="Sole-source suppliers (only active company in the category): all kept "
            "by rule R2").font = font(bold=True)
    write_df(ws, ss, r, {"Spend, 3 yrs ($)": USD})
    widths(ws, [56, 16, 16, 14, 14, 14, 14, 14, 12, 12, 18, 18])

if Path("outputs/step4_duplicate_review_sample.csv").exists():
    ws = wb.create_sheet("4_Duplicate_Check")
    title(ws, "Step 4d — Precision of duplicate matching at the 0.88 threshold",
          "Fixed random sample (seed 42) of accepted pairs, judged by hand from name, ZIP code, main category "
          "and department. Uncertain pairs counted as wrong.")
    pr = pd.read_csv("outputs/step4d_duplicate_precision.csv")
    pr.columns = ["Pair type", "Sampled", "Correct", "Precision", "Precision if uncertain = correct",
                  "Accepted pairs in data"]
    end = write_df(ws, pr, 4, {"Sampled": NUM, "Correct": NUM, "Precision": PCT,
                               "Precision if uncertain = correct": PCT, "Accepted pairs in data": NUM})
    rv = pd.read_csv("outputs/step4_duplicate_review_sample.csv")
    rv = rv[["id", "kind", "similarity", "name_a", "name_b", "ctx_a", "ctx_b", "verdict", "uncertain", "note"]]
    rv.columns = ["#", "Type", "Similarity", "Name A", "Name B", "A: ZIP | category | department",
                  "B: ZIP | category | department", "Verdict", "Uncertain", "Note"]
    write_df(ws, rv, end + 3, {"Similarity": "0.000"})
    widths(ws, [24, 9, 10, 40, 40, 52, 52, 10, 10, 60])

if Path("outputs/step5b_retender_sensitivity.csv").exists():
    ws = wb.create_sheet("5b_Retender_Sensitivity")
    title(ws, "Step 5b — How much of the case rests on the re-tender assumption?",
          "J&J run-rate savings = measured savings + re-tender rate x kept Leverage spend. Blue = input.")
    sc = pd.read_csv("outputs/step5_savings_scenarios.csv").set_index("scenario")
    inp = pd.read_csv("outputs/step5_model_inputs.csv").set_index("input").value
    kept_lev = inp["Kept spend per year, Leverage categories ($M)"] * 35000 / inp["CA addressable spend per year ($M)"]
    rows_in = [("Measured savings, base: price gap + fewer suppliers ($M/yr)", sc.loc["Base", "jnj_price_m"] + sc.loc["Base", "jnj_process_m"], USD_M),
               ("Kept Leverage spend at J&J scale ($M/yr)", kept_lev, USD_M),
               ("Goal ($M/yr)", 300, USD_M)]
    header(ws, 4, ["Input", "Value"])
    for i, (lab, v, fmt) in enumerate(rows_in, start=5):
        ws.cell(row=i, column=1, value=lab).font = font()
        c = ws.cell(row=i, column=2, value=round(float(v), 3))
        c.font, c.number_format = font(color=BLUE), fmt
    ws.cell(row=9, column=1, value="Break-even re-tender rate for the goal").font = font(bold=True)
    c = ws.cell(row=9, column=2, value="=(B7-B5)/B6")
    c.font, c.number_format, c.fill = font(bold=True), '0.00%', KEY_FILL
    header(ws, 11, ["Re-tender rate", "J&J run-rate savings ($M/yr)", "Meets goal?"])
    for i, r_ in enumerate([x / 1000 for x in range(0, 31, 1)], start=12):
        ws.cell(row=i, column=1, value=r_).number_format = '0.0%'
        ws.cell(row=i, column=2, value=f"=$B$5+A{i}*$B$6").number_format = USD_M
        ws.cell(row=i, column=3, value=f'=IF(B{i}>=$B$7,"Yes","No")')
    widths(ws, [58, 28, 14])


# ---------------------------------------------------------------- Step 5c / 8 / 9
if Path("outputs/step5c_value_vs_cuts.csv").exists():
    ws = wb.create_sheet("5c_Value_vs_Cuts")
    title(ws, "Step 5c — Small cut, big impact: savings vs number of suppliers cut (J&J scale, base savings rates)",
          "Every cut has to be justified, so compare what each tier adds. J&J baseline: 80,000 suppliers.")
    v = pd.read_csv("outputs/step5c_value_vs_cuts.csv")
    v = v[["tier", "reduction", "jnj_suppliers_cut", "jnj_after", "base_rate_price_m", "base_rate_process_m",
           "base_rate_retender_m", "base_rate_run_rate_m", "base_rate_one_time_m", "share_of_base_value",
           "savings_per_supplier_cut_k", "suppliers_holding_80pct_of_kept_leverage_spend"]]
    v.columns = ["Tier", "Share of active suppliers cut", "J&J suppliers cut", "J&J suppliers after", "Price gap ($M/yr)",
                 "Fewer suppliers ($M/yr)", "Re-tender, assumed ($M/yr)", "Run-rate ($M/yr)", "One-time ($M)",
                 "Share of base-case value", "Savings per supplier cut ($k/yr)", "Suppliers holding 80% of kept Leverage spend"]
    write_df(ws, v, 4, {"Share of active suppliers cut": PCT, "J&J suppliers cut": NUM, "J&J suppliers after": NUM,
                        "Price gap ($M/yr)": USD_M, "Fewer suppliers ($M/yr)": USD_M, "Re-tender, assumed ($M/yr)": USD_M,
                        "Run-rate ($M/yr)": USD_M, "One-time ($M)": USD_M, "Share of base-case value": PCT,
                        "Savings per supplier cut ($k/yr)": '$#,##0.0', "Suppliers holding 80% of kept Leverage spend": NUM})
    widths(ws, [14, 14, 14, 14, 14, 14, 16, 14, 12, 14, 16, 18])

if Path("outputs/step8_phase2_summary.csv").exists():
    ws = wb.create_sheet("8_Phase2")
    title(ws, "Phase 2 — Category target count from ~39,000 suppliers",
          "Per category: fewest suppliers covering 90% of spend; at least 2 in Strategic/Bottleneck; strategic, "
          "sole-source and small/diverse suppliers always kept.")
    ps = pd.read_csv("outputs/step8_phase2_summary.csv")
    ps.columns = ["Metric", "Value"]
    end = write_df(ws, ps, 4)
    for r in range(5, end + 1):
        c = ws.cell(row=r, column=2)
        c.number_format = PCT if isinstance(c.value, float) and c.value < 1 else '#,##0.0'
    pq = pd.read_csv("outputs/step8_phase2_by_quadrant.csv")
    pq.columns = ["Quadrant", "Suppliers before", "Suppliers after", "Reduction"]
    end = write_df(ws, pq, end + 3, {"Suppliers before": NUM, "Suppliers after": NUM, "Reduction": PCT})
    sv = pd.read_csv("outputs/step8_phase2_coverage_sensitivity.csv")
    sv.columns = ["Spend coverage target", "California suppliers kept", "Reduction", "J&J suppliers after"]
    write_df(ws, sv, end + 3, {"Spend coverage target": PCT, "California suppliers kept": NUM, "Reduction": PCT,
                               "J&J suppliers after": NUM})
    widths(ws, [62, 18, 16, 18])

if Path("outputs/step9_jnj_timeline_backtest.csv").exists():
    ws = wb.create_sheet("9_JnJ_Timeline")
    title(ws, "Back-test — J&J's supplier count over time and the Kenvue spin-off",
          "80,000 (State of Flux; ~78,000 in a 2021 StoryMap) to ~39,000 (FY2025). Kenvue's share is estimated.")
    rep_rows = pd.DataFrame([
        ("After 2014 (undated)", 80000, "State of Flux case study"),
        ("2021", 78000, "ArcGIS StoryMap, individual author, no source given"),
        ("2023", 38500, "2023 Health for Humanity Report"),
        ("2024", 29000, "2024 Health for Humanity Report (after Kenvue)"),
        ("FY2025", 39000, "Global Modern Slavery Statement 2025 (new counting method)"),
    ], columns=["When", "Suppliers reported", "Source"])
    end = write_df(ws, rep_rows, 4, {"Suppliers reported": NUM})
    bt = pd.read_csv("outputs/step9_jnj_timeline_backtest.csv")
    bt.columns = ["Kenvue case", "Kenvue share of suppliers", "Kenvue suppliers", "J&J base after Kenvue",
                  "Reported now", "J&J's own reduction", "Our model: conservative", "Our model: base",
                  "Our model: aggressive", "Phase 2 (90% coverage)", "Phase 2 low", "Phase 2 high"]
    write_df(ws, bt, end + 3, {**{c: NUM for c in bt.columns[2:5]}, "Kenvue share of suppliers": PCT,
                               "J&J's own reduction": PCT, **{c: NUM for c in bt.columns[6:]}})
    widths(ws, [22, 18, 52, 18, 14, 16, 16, 16, 16, 18, 14, 14])

# ---------------------------------------------------------------- Files and sources
REPO = "https://github.com/hemant-singla/jnj-supplier-rationalization/blob/main/"
ws = wb.create_sheet("Files_and_Sources", 1)
title(ws, "Files and sources used in this workbook",
      "Every number traces to one of these. Code links point to the project's GitHub repository.")
files = [
    ("Data source", "All", "State of California DGS — Purchase Order Data FY2012-15 (CSV, 157 MB, public)",
     "https://catalog.data.gov/dataset/purchase-order-data", "Real purchase orders: suppliers, amounts, UNSPSC categories"),
    ("Data source", "0", "J&J 2024 Health for Humanity Report",
     "https://healthforhumanityreport.jnj.com/2024/_assets/downloads/johnson-johnson-2024-health-for-humanity-report.pdf?h=EDX6Vl9g",
     "29,000+ suppliers (2024)"),
    ("Data source", "0", "J&J 2023 Health for Humanity Report",
     "https://healthforhumanityreport.jnj.com/2023/_assets/downloads/johnson-johnson-2023-health-for-humanity-report.pdf?h=Ka9OvM1t",
     "38,500+ suppliers (2023)"),
    ("Data source", "0", "J&J careers: Director, Procurement Global Process Owner",
     "https://www.careers.jnj.com/en/jobs/r-088156/director-procurement-global-process-owner/",
     "~$35B spend across ~35,000 suppliers"),
    ("Data source", "0", "J&J: DEI in our supply base (2022)",
     "https://belong.jnj.com/2022/driving-innovation/dei-in-our-supply-base.html", "$3B Tier 1 diverse spend (2022)"),
    ("Data source", "—", "SEC EDGAR full-text search (10-K / 20-F, 2023-2026)", "https://efts.sec.gov/LATEST/search-index",
     "Real J&J suppliers verified from their 10-K filings (docs/07_real_jnj_suppliers.md)"),
    ("Code", "0", "model/business_case_model.py", REPO + "model/business_case_model.py",
     "Business case model (same logic as tab 0_Business_Case)"),
    ("Code", "1", "scripts/01_clean_load.py", REPO + "scripts/01_clean_load.py", "Cleans the CSV, applies exclusion rules, loads SQLite"),
    ("SQL", "2", "sql/01_pareto_abc.sql", REPO + "sql/01_pareto_abc.sql", "Pareto / ABC segment per supplier"),
    ("SQL", "2", "sql/02_category_fragmentation.sql", REPO + "sql/02_category_fragmentation.sql", "Category cube"),
    ("Code", "2", "scripts/02_spend_cube.py", REPO + "scripts/02_spend_cube.py", "Runs the step 2 SQL, saves outputs"),
    ("Code", "3", "scripts/03_kraljic.py", REPO + "scripts/03_kraljic.py", "Kraljic scoring and 2x2 chart"),
    ("Code", "4", "scripts/04a_duplicates.py", REPO + "scripts/04a_duplicates.py", "TF-IDF duplicate detection"),
    ("Code", "4", "scripts/04b_decisions.py", REPO + "scripts/04b_decisions.py", "Scorecard and decision rules"),
    ("Code", "4", "scripts/04c_concentration.py", REPO + "scripts/04c_concentration.py", "Supplier concentration before vs after; sole-source list"),
    ("Code", "4", "scripts/04d_duplicate_precision.py", REPO + "scripts/04d_duplicate_precision.py", "Precision of duplicate matching from a hand-checked sample"),
    ("Code", "5", "scripts/05_savings.py", REPO + "scripts/05_savings.py", "Price-gap test and savings model"),
    ("Code", "5", "scripts/05b_retender_sensitivity.py", REPO + "scripts/05b_retender_sensitivity.py", "Savings vs re-tender rate; break-even"),
    ("Code", "5", "scripts/05c_value_vs_cuts.py", REPO + "scripts/05c_value_vs_cuts.py", "Small cut, big impact: savings vs suppliers cut"),
    ("Code", "8", "scripts/08_phase2_category_targets.py", REPO + "scripts/08_phase2_category_targets.py", "Phase 2 category target model"),
    ("Code", "9", "scripts/09_jnj_timeline_backtest.py", REPO + "scripts/09_jnj_timeline_backtest.py", "J&J timeline, Kenvue estimate, back-test"),
    ("Write-up", "8-9", "docs/08_jnj_timeline_and_phase2.md", REPO + "docs/08_jnj_timeline_and_phase2.md", "80,000 to ~39,000, Kenvue, small cut big impact, Phase 2"),
    ("Data source", "9", "State of Flux: J&J case study", "https://stateofflux.co.uk/case-studies/johnson-and-johnson", "80,000 suppliers (undated)"),
    ("Data source", "9", "ArcGIS StoryMap: J&J Supply Chain (2021)", "https://storymaps.arcgis.com/stories/efd9fb73ac4f4e8082036d6296aea5a2", "78,000+ suppliers (individual author)"),
    ("Data source", "9", "J&J Global Modern Slavery Statement 2025", "https://www.jnj.com/innovativemedicine/uk/global-modern-slavery-act-statement", "~39,000 suppliers (FY2025)"),
    ("Data source", "9", "J&J Q4 and full-year 2022 results", "https://www.businesswire.com/news/home/20230124005275/en", "Consumer Health $14,953M of $94,943M sales (Kenvue share estimate)"),
    ("Data source", "9", "NJBIZ: Kenvue completes separation", "https://njbiz.com/kenvue-completes-separation-from-jj", "Separation completed 23 Aug 2023"),
    ("Tests", "All", "tests/", REPO + "tests", "Checks: no strategic or sole-source exits, a reason code on every row, totals match the raw data"),
    ("Code", "All", "scripts/build_outputs_workbook.py", REPO + "scripts/build_outputs_workbook.py", "Builds this workbook"),
    ("Code", "1-4", "scripts/build_data_workbook.py", REPO + "scripts/build_data_workbook.py",
     "Builds the raw-data workbook (every raw row traced through steps 1-4)"),
    ("Output", "1-4", "outputs/JnJ_Data_Raw_to_Decision_part1of5.xlsx ... part5of5 (100 MB in total, not in GitHub)",
     REPO + "scripts/build_data_workbook.py",
     "All 344,504 raw rows as downloaded + status, exclusion reason, segment, quadrant, decision per row"),
    ("Output", "2", "outputs/step2_*.csv", REPO + "outputs", "Segment summary, Pareto curve, fragmented categories"),
    ("Output", "3", "outputs/step3_kraljic_*.csv; docs/images/kraljic_matrix.png", REPO + "outputs", "Quadrant per category; chart"),
    ("Output", "4", "outputs/step4_*.csv, step4c_*.csv, step4d_*.csv", REPO + "outputs",
     "Supplier decisions, decisions by rule/segment, headline, duplicate pairs, concentration, duplicate precision"),
    ("Output", "5", "outputs/step5_*.csv, step5b_*.csv; docs/images/savings_bridge_real.png, retender_sensitivity.png", REPO + "outputs",
     "Price-gap items, measured gap by quadrant, model inputs, scenarios, sensitivity; charts"),
    ("Write-up", "0", "docs/01_business_case.md", REPO + "docs/01_business_case.md", "Business case"),
    ("Write-up", "—", "docs/02_methodology.md", REPO + "docs/02_methodology.md", "Goal and methods"),
    ("Write-up", "2", "docs/03_step2_spend_cube.md", REPO + "docs/03_step2_spend_cube.md", "Step 2 findings"),
    ("Write-up", "3", "docs/04_step3_kraljic.md", REPO + "docs/04_step3_kraljic.md", "Step 3 findings"),
    ("Write-up", "4", "docs/05_step4_decisions.md", REPO + "docs/05_step4_decisions.md", "Step 4 findings"),
    ("Write-up", "5", "docs/06_step5_savings.md", REPO + "docs/06_step5_savings.md", "Step 5 findings"),
    ("Code", "6", "scripts/06_powerbi_data.py; scripts/07_build_pbip.py", REPO + "powerbi/README.md", "Power BI data file and project generator"),
    ("Output", "6", "powerbi/JnJ_Supplier_Dashboard (Power BI project)", REPO + "powerbi", "4-page dashboard"),
    ("Write-up", "6b", "docs/07_real_jnj_suppliers.md", REPO + "docs/07_real_jnj_suppliers.md",
     "Verified real J&J suppliers from SEC 10-K filings: Alkermes, MeiraGTx, Mersana, Emergent (former)"),
    ("Write-up", "7", "docs/agreements/", REPO + "docs/agreements/README.md", "MOU, MSA, bottleneck supply, catalog terms, transition notices"),
    ("Database", "1-4", "data/procurement.db (SQLite, rebuilt locally, not in git)", REPO + "data/README.md",
     "Tables: spend_lines, excluded_lines, supplier_segments, category_cube, supplier_duplicates, supplier_decisions"),
]
header(ws, 4, ["Type", "Step", "File / source", "Link", "What it provides"])
for i, rec in enumerate(files, start=5):
    for c, v in enumerate(rec, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font = font()
        if c == 4:
            cell.hyperlink = v
            cell.font = Font(name=F, size=10, color="0563C1", underline="single")
widths(ws, [11, 6, 52, 60, 70])

ws = wb.create_sheet("SQL_Code")
title(ws, "SQL used in step 2 (runs on SQLite)", "Full text of the queries; also in the repo's sql/ folder.")
r = 4
for fname in ["sql/01_pareto_abc.sql", "sql/02_category_fragmentation.sql"]:
    ws.cell(row=r, column=1, value=fname).font = font(bold=True)
    r += 1
    for ln in Path(fname).read_text().splitlines():
        ws.cell(row=r, column=1, value=ln).font = Font(name="Consolas", size=9)
        r += 1
    r += 2
widths(ws, [130])

# keep tabs in step order
for name, after in [("4_Concentration", "4_Supplier_Decisions"), ("4_Duplicate_Check", "4_Concentration")]:
    if name in wb.sheetnames and after in wb.sheetnames:
        wb.move_sheet(name, offset=wb.sheetnames.index(after) + 1 - wb.sheetnames.index(name))
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("saved", OUT)
